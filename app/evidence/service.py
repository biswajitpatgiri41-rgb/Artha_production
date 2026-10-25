"""Main evidence service orchestrating upload, parsing, indexing, and retrieval."""

from __future__ import annotations

import asyncio
import io
import json
import threading
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd

from app.core.config import get_evidence_limits, get_settings
from app.core.security import detect_prompt_injection
from app.evidence.schemas import (
    ALLOWED_MIME_TYPES,
    EvidenceDocument,
    EvidenceFormat,
    EvidenceLifecycle,
    EvidenceMetadata,
    EvidenceStatus,
    RetrievedEvidence,
    RetrievedEvidenceLocation,
    SearchRequest,
    SearchResponse,
    StructuredJSONResult,
    StructuredTableResult,
    UploadResponse,
)
from app.evidence.storage import EvidenceStorage, get_storage
from app.evidence.parsers import FormatRouter, get_format_router, parse_file
from app.evidence.chunking import chunk_document, ChunkConfig
from app.evidence.embeddings import get_embedding_provider
from app.evidence.vector_store import get_vector_store
from app.evidence.retrieval import retrieve_evidence
from app.evidence.events import (
    emit_evidence_upload_started,
    emit_evidence_uploaded,
    emit_evidence_parsing_started,
    emit_evidence_parsing_completed,
    emit_evidence_processing_failed,
    emit_evidence_ready,
    emit_evidence_indexed,
    emit_retrieval_started,
    emit_retrieval_completed,
    emit_evidence_validation_started,
    emit_evidence_validation_completed,
    emit_evidence_deleted,
)
class StructuredTableService:
    """Service for querying structured table evidence (CSV, XLSX)."""

    def __init__(self, storage: EvidenceStorage):
        self.storage = storage

    def _load_dataframe(self, run_id: str, evidence_id: str, sheet: Optional[str] = None) -> pd.DataFrame:
        metadata = self._get_metadata(run_id, evidence_id)
        if not metadata:
            raise ValueError(f"Evidence {evidence_id} not found")
        file_path = self.storage.get_file_path(run_id, metadata.stored_filename)
        if metadata.format == EvidenceFormat.CSV:
            return pd.read_csv(file_path, low_memory=False, nrows=get_evidence_limits()["max_csv_rows"] + 1)
        elif metadata.format == EvidenceFormat.XLSX:
            sheet_name = sheet or (metadata.sheet_names[0] if metadata.sheet_names else 0)
            if metadata.sheet_names and sheet_name not in metadata.sheet_names:
                raise ValueError("Unknown workbook sheet")
            from openpyxl import load_workbook
            workbook = load_workbook(file_path, read_only=True, data_only=True, keep_links=False)
            try:
                worksheet = workbook[sheet_name]
                values = worksheet.iter_rows(values_only=True)
                header = next(values, ())
                columns = [str(value) if value is not None else f"column_{index + 1}" for index, value in enumerate(header)]
                rows = []
                for row in values:
                    if len(rows) >= get_evidence_limits()["max_csv_rows"] + 1:
                        break
                    rows.append(row)
                return pd.DataFrame(rows, columns=columns)
            finally:
                workbook.close()
        else:
            raise ValueError(f"Format {metadata.format} not supported for table queries")

    def _get_metadata(self, run_id: str, evidence_id: str) -> Optional[EvidenceMetadata]:
        return self.storage.load_metadata(run_id, evidence_id)

    def list_columns(self, run_id: str, evidence_id: str, sheet: Optional[str] = None) -> List[str]:
        df = self._load_dataframe(run_id, evidence_id, sheet)
        return df.columns.tolist()

    def get_schema(self, run_id: str, evidence_id: str, sheet: Optional[str] = None) -> Dict[str, str]:
        df = self._load_dataframe(run_id, evidence_id, sheet)
        return {col: str(dtype) for col, dtype in df.dtypes.items()}

    def get_row_count(self, run_id: str, evidence_id: str, sheet: Optional[str] = None) -> int:
        df = self._load_dataframe(run_id, evidence_id, sheet)
        return len(df)

    def filter_rows(self, run_id: str, evidence_id: str, conditions: Dict[str, Any], sheet: Optional[str] = None, limit: int = 100) -> StructuredTableResult:
        df = self._load_dataframe(run_id, evidence_id, sheet)
        mask = pd.Series([True] * len(df), index=df.index)
        for col, value in conditions.items():
            if col not in df.columns:
                raise ValueError(f"Column {col} not found")
            mask = mask & (df[col] == value)
        filtered = df[mask].head(limit)
        metadata = self._get_metadata(run_id, evidence_id)
        data = json.loads(filtered.to_json(orient="records", date_format="iso"))
        return StructuredTableResult(
            source_file=metadata.original_filename if metadata else "unknown",
            sheet=sheet,
            rows_involved=[int(index) + 2 for index in filtered.index],
            columns_involved=filtered.columns.tolist(),
            data=data,
            evidence_reference=f"{evidence_id}:{sheet or 'default'}",
        )

    def describe_numeric_columns(self, run_id: str, evidence_id: str, sheet: Optional[str] = None) -> Dict[str, Dict[str, float]]:
        df = self._load_dataframe(run_id, evidence_id, sheet)
        numeric = df.select_dtypes(include=["number"])
        return numeric.describe().to_dict()

    def lookup_value(self, run_id: str, evidence_id: str, column: str, value: Any, sheet: Optional[str] = None) -> StructuredTableResult:
        return self.filter_rows(run_id, evidence_id, {column: value}, sheet)

    def query_date_range(self, run_id: str, evidence_id: str, date_column: str, start_date: str, end_date: str, sheet: Optional[str] = None) -> StructuredTableResult:
        df = self._load_dataframe(run_id, evidence_id, sheet)
        if date_column not in df.columns:
            raise ValueError(f"Column {date_column} not found")
        df[date_column] = pd.to_datetime(df[date_column], errors="coerce")
        start = pd.to_datetime(start_date, errors="raise")
        end = pd.to_datetime(end_date, errors="raise")
        mask = (df[date_column] >= start) & (df[date_column] <= end)
        filtered = df[mask]
        metadata = self._get_metadata(run_id, evidence_id)
        data = json.loads(filtered.to_json(orient="records", date_format="iso"))
        return StructuredTableResult(
            source_file=metadata.original_filename if metadata else "unknown",
            sheet=sheet,
            rows_involved=[int(index) + 2 for index in filtered.index],
            columns_involved=filtered.columns.tolist(),
            data=data,
            evidence_reference=f"{evidence_id}:{sheet or 'default'}",
        )
class StructuredJSONService:
    """Service for querying structured JSON evidence."""

    def __init__(self, storage: EvidenceStorage):
        self.storage = storage

    def _load_json(self, run_id: str, evidence_id: str) -> Any:
        metadata = self._get_metadata(run_id, evidence_id)
        if not metadata:
            raise ValueError(f"Evidence {evidence_id} not found")
        file_path = self.storage.get_file_path(run_id, metadata.stored_filename)
        return json.loads(file_path.read_text(encoding="utf-8"))

    def _get_metadata(self, run_id: str, evidence_id: str) -> Optional[EvidenceMetadata]:
        return self.storage.load_metadata(run_id, evidence_id)

    def list_keys(self, run_id: str, evidence_id: str) -> List[str]:
        data = self._load_json(run_id, evidence_id)
        return list(data.keys()) if isinstance(data, dict) else []

    def get_path(self, run_id: str, evidence_id: str, path: str) -> Any:
        data = self._load_json(run_id, evidence_id)
        keys = path.split(".")
        current = data
        for key in keys:
            if isinstance(current, dict):
                current = current.get(key)
            elif isinstance(current, list) and key.isdigit():
                current = current[int(key)]
            else:
                return None
        return current

    def search_values(self, run_id: str, evidence_id: str, search_term: str) -> List[Dict[str, Any]]:
        data = self._load_json(run_id, evidence_id)
        results = []
        visited = 0
        def _search(obj: Any, path: str = ""):
            nonlocal visited
            visited += 1
            if visited > 200_000 or len(results) >= 1000:
                return
            if isinstance(obj, dict):
                for k, v in obj.items():
                    if search_term.casefold() in str(k).casefold():
                        results.append({"path": f"{path}.{k}" if path else str(k), "value": v if isinstance(v, (str, int, float, bool)) or v is None else "[structured value]"})
                    _search(v, f"{path}.{k}" if path else k)
            elif isinstance(obj, list):
                for i, v in enumerate(obj):
                    _search(v, f"{path}[{i}]")
            elif search_term.casefold() in str(obj).casefold():
                results.append({"path": path, "value": obj if isinstance(obj, (str, int, float, bool)) or obj is None else "[structured value]"})
        _search(data)
        return results

    def extract_records(self, run_id: str, evidence_id: str, path: str) -> StructuredJSONResult:
        current = self.get_path(run_id, evidence_id, path)
        metadata = self._get_metadata(run_id, evidence_id)
        records = current[:1000] if isinstance(current, list) else []
        return StructuredJSONResult(
            source_file=metadata.original_filename if metadata else "unknown", path=path,
            keys_involved=path.split("."), data=records, evidence_reference=evidence_id
        )
class EvidenceService:
    """Main service for evidence operations."""

    def __init__(self):
        self.storage = get_storage()
        self.format_router = get_format_router()
        self.vector_store = get_vector_store()
        self.embedding_provider = get_embedding_provider()
        self.table_service = StructuredTableService(self.storage)
        self.json_service = StructuredJSONService(self.storage)
        self._metadata_cache: Dict[str, EvidenceMetadata] = {}
        self._lifecycle_lock = threading.RLock()
        self._deleting: set[str] = set()

    async def upload_file(
        self, 
        run_id: str, 
        original_filename: str, 
        mime_type: str, 
        file_content: bytes,
        lifecycle: EvidenceLifecycle = EvidenceLifecycle.TEMPORARY
    ) -> UploadResponse:
        file_size = len(file_content)
        valid, error = self.storage.validate_file(original_filename, mime_type, file_size)
        if not valid:
            raise ValueError(error)
        
        # Emit upload started event
        emit_evidence_upload_started(run_id=run_id, filename=original_filename, size=file_size, lifecycle=lifecycle.value)
        
        metadata, file_path = self.storage.save_upload(run_id, original_filename, mime_type, io.BytesIO(file_content), file_size)
        metadata.lifecycle = lifecycle
        cache_key = f"{run_id}:{metadata.evidence_id}"
        self._metadata_cache[cache_key] = metadata
        self.storage.save_metadata(metadata)
        
        # Emit upload completed event
        emit_evidence_uploaded(
            run_id=run_id, 
            evidence_id=metadata.evidence_id, 
            filename=original_filename, 
            format=metadata.format.value,
            size=file_size,
            lifecycle=lifecycle.value
        )
        
        asyncio.create_task(asyncio.to_thread(self._process_file, run_id, metadata, file_path))
        return UploadResponse(
            evidence_id=metadata.evidence_id, 
            run_id=run_id, 
            filename=original_filename, 
            format=metadata.format, 
            size=file_size, 
            status=metadata.status, 
            message="Upload successful. Processing started."
        )

    def _process_file(self, run_id: str, metadata: EvidenceMetadata, file_path: Path) -> None:
        cache_key = f"{run_id}:{metadata.evidence_id}"
        try:
            with self._lifecycle_lock:
                if cache_key in self._deleting or not self.storage.load_metadata(run_id, metadata.evidence_id):
                    return
            # Emit parsing started event
            emit_evidence_parsing_started(run_id=run_id, evidence_id=metadata.evidence_id, format=metadata.format.value)
            
            doc = parse_file(file_path, metadata)
            if doc.text and detect_prompt_injection(doc.text)[0]:
                metadata.parser_warnings.append(
                    "Potential instruction-like text detected; retained only as untrusted evidence data."
                )
            with self._lifecycle_lock:
                if cache_key in self._deleting or not self.storage.load_metadata(run_id, metadata.evidence_id):
                    return
                self.storage.save_metadata(metadata)
            emit_evidence_parsing_completed(run_id, metadata.evidence_id, metadata.format.value, len(metadata.parser_warnings))
            
            if metadata.format in (EvidenceFormat.PDF, EvidenceFormat.TXT, EvidenceFormat.MD):
                chunks = chunk_document(doc)
                if chunks:
                    metadata.chunk_count = len(chunks)
                    with self._lifecycle_lock:
                        if cache_key in self._deleting or not self.storage.load_metadata(run_id, metadata.evidence_id):
                            return
                        metadata.status = EvidenceStatus.INDEXING
                        self.storage.save_metadata(metadata)
                    texts = [c["text"] for c in chunks]
                    embeddings = self.embedding_provider.embed(texts)
                    with self._lifecycle_lock:
                        if cache_key in self._deleting or not self.storage.load_metadata(run_id, metadata.evidence_id):
                            return
                        self.vector_store.add_chunks(run_id, chunks, embeddings)
                    metadata.indexed_chunk_count = len(chunks)
                    
                    # Emit indexed event
                    emit_evidence_indexed(
                        run_id=run_id, 
                        evidence_id=metadata.evidence_id, 
                        chunk_count=len(chunks),
                        format=metadata.format.value
                    )
            with self._lifecycle_lock:
                if cache_key in self._deleting or not self.storage.load_metadata(run_id, metadata.evidence_id):
                    return
                metadata.status = EvidenceStatus.READY
                self.storage.save_metadata(metadata)
            emit_evidence_ready(run_id, metadata.evidence_id, metadata.format.value, metadata.indexed_chunk_count or 0)
        except Exception as e:
            with self._lifecycle_lock:
                if cache_key in self._deleting or not self.storage.load_metadata(run_id, metadata.evidence_id):
                    return
                metadata.status = EvidenceStatus.FAILED
                metadata.error_message = f"Evidence processing failed ({type(e).__name__}). See server logs for details."
                self.storage.save_metadata(metadata)
            emit_evidence_processing_failed(run_id, metadata.evidence_id, metadata.format.value, type(e).__name__)

    def get_metadata(self, run_id: str, evidence_id: str) -> Optional[EvidenceMetadata]:
        cache_key = f"{run_id}:{evidence_id}"
        metadata = self._metadata_cache.get(cache_key)
        if metadata is None:
            metadata = self.storage.load_metadata(run_id, evidence_id)
            if metadata is not None:
                self._metadata_cache[cache_key] = metadata
        return metadata

    def list_run_evidence(self, run_id: str) -> List[EvidenceMetadata]:
        self.cleanup_expired_evidence()
        persisted = self.storage.list_metadata(run_id)
        for metadata in persisted:
            self._metadata_cache[f"{run_id}:{metadata.evidence_id}"] = metadata
        return persisted

    def delete_evidence(self, run_id: str, evidence_id: str) -> bool:
        """Delete evidence and all associated data.
        
        Removes:
        - uploaded file
        - parsed content (chunks from vector store)
        - vector entries for this evidence
        - metadata references
        
        Does NOT corrupt completed audit records.
        """
        cache_key = f"{run_id}:{evidence_id}"
        with self._lifecycle_lock:
            metadata = self.get_metadata(run_id, evidence_id)
            if not metadata:
                return False
            # Mark first so an in-flight parser cannot recreate metadata or
            # index a vector after this removal acquires the lifecycle lock.
            self._deleting.add(cache_key)
            vector_error = None
            try:
                self.vector_store.delete_evidence(run_id, evidence_id)
            except Exception as exc:  # Continue removal; stale vectors are run/id filtered.
                vector_error = exc
            self.storage.delete_file(run_id, metadata.stored_filename)
            self.storage.delete_metadata(run_id, evidence_id)
            self._metadata_cache.pop(cache_key, None)
            self._deleting.discard(cache_key)
            if vector_error is not None:
                raise RuntimeError("Evidence files were removed, but vector cleanup failed") from vector_error
            emit_evidence_deleted(run_id, evidence_id, metadata.lifecycle.value)
            return True

    def cleanup_expired_evidence(self) -> int:
        """Clean up temporary evidence that has exceeded retention period.
        
        Returns the number of evidence items cleaned up.
        """
        limits = get_evidence_limits()
        retention_hours = limits["evidence_retention_hours"]
        cutoff_time = datetime.now(timezone.utc) - timedelta(hours=retention_hours)
        
        cleaned_count = 0
        for metadata in self.storage.list_all_metadata():
            if metadata.lifecycle == EvidenceLifecycle.TEMPORARY:
                uploaded_at = metadata.uploaded_at
                if uploaded_at.tzinfo is None:
                    uploaded_at = uploaded_at.replace(tzinfo=timezone.utc)
                if uploaded_at < cutoff_time:
                    if self.delete_evidence(metadata.run_id, metadata.evidence_id):
                        cleaned_count += 1
        
        return cleaned_count

    async def search(self, request: SearchRequest) -> SearchResponse:
        self.cleanup_expired_evidence()
        # Emit retrieval started event
        emit_retrieval_started(run_id=request.run_id, query=request.query, top_k=request.top_k)
        
        response = await retrieve_evidence(request)
        
        # Emit retrieval completed event
        emit_retrieval_completed(
            run_id=request.run_id, 
            query=request.query, 
            result_count=len(response.results),
            top_k=request.top_k
        )
        
        return response

    def validate_evidence(self, run_id: str, evidence_id: str) -> Dict[str, Any]:
        """Validate evidence content for consistency and quality."""
        metadata = self.get_metadata(run_id, evidence_id)
        if not metadata:
            return {"valid": False, "error": "Evidence not found"}
        
        # Emit validation started event
        emit_evidence_validation_started(run_id=run_id, evidence_id=evidence_id, format=metadata.format.value)
        
        # Perform validation checks
        validation_result = {
            "valid": True,
            "evidence_id": evidence_id,
            "format": metadata.format.value,
            "checks": {}
        }
        
        # Check if evidence is ready
        validation_result["checks"]["status_ready"] = metadata.status == EvidenceStatus.READY
        
        # Check if vector entries exist for text formats
        if metadata.format in (EvidenceFormat.PDF, EvidenceFormat.TXT, EvidenceFormat.MD):
            validation_result["checks"]["has_chunks"] = metadata.chunk_count is not None and metadata.chunk_count > 0
            validation_result["checks"]["indexed_chunks"] = metadata.indexed_chunk_count is not None and metadata.indexed_chunk_count > 0
            validation_result["checks"]["all_indexed"] = (
                metadata.chunk_count is not None 
                and metadata.indexed_chunk_count is not None
                and metadata.chunk_count == metadata.indexed_chunk_count
            )
        
        # Check structured data
        if metadata.format in (EvidenceFormat.CSV, EvidenceFormat.XLSX):
            validation_result["checks"]["has_tables"] = metadata.row_count is not None and metadata.row_count > 0
            validation_result["checks"]["has_columns"] = metadata.columns is not None and len(metadata.columns) > 0
        
        if metadata.format == EvidenceFormat.JSON:
            validation_result["checks"]["has_data"] = metadata.top_level_keys is not None and len(metadata.top_level_keys) > 0
        
        # Overall validity
        validation_result["valid"] = all(validation_result["checks"].values())
        
        # Emit validation completed event
        emit_evidence_validation_completed(
            run_id=run_id, 
            evidence_id=evidence_id, 
            valid=validation_result["valid"],
            checks=validation_result["checks"]
        )
        
        return validation_result

    def query_table(self, run_id: str, evidence_id: str, operation: str, **kwargs) -> StructuredTableResult:
        ops = {
            "list_columns": self.table_service.list_columns,
            "get_schema": self.table_service.get_schema,
            "get_row_count": self.table_service.get_row_count,
            "filter_rows": self.table_service.filter_rows,
            "describe_numeric": self.table_service.describe_numeric_columns,
            "lookup_value": self.table_service.lookup_value,
            "query_date_range": self.table_service.query_date_range,
        }
        if operation not in ops:
            raise ValueError(f"Unknown table operation: {operation}")
        result = ops[operation](run_id, evidence_id, **kwargs)
        if isinstance(result, StructuredTableResult):
            return result
        metadata = self.storage.load_metadata(run_id, evidence_id)
        data = result if isinstance(result, list) and all(isinstance(row, dict) for row in result) else [{"result": result}]
        return StructuredTableResult(
            source_file=metadata.original_filename if metadata else "unknown",
            sheet=kwargs.get("sheet"), rows_involved=[],
            columns_involved=list(result.keys()) if isinstance(result, dict) else [],
            data=data, evidence_reference=f"{evidence_id}:{kwargs.get('sheet') or 'default'}",
        )

    def query_json(self, run_id: str, evidence_id: str, operation: str, **kwargs) -> StructuredJSONResult:
        ops = {
            "list_keys": self.json_service.list_keys,
            "get_path": self.json_service.get_path,
            "search_values": self.json_service.search_values,
            "extract_records": self.json_service.extract_records,
        }
        if operation not in ops:
            raise ValueError(f"Unknown JSON operation: {operation}")
        result = ops[operation](run_id, evidence_id, **kwargs)
        if isinstance(result, StructuredJSONResult):
            return result
        metadata = self.storage.load_metadata(run_id, evidence_id)
        path = kwargs.get("path", "$")
        keys = path.split(".") if isinstance(path, str) else []
        data = result if isinstance(result, (dict, list, str, int, float, bool)) or result is None else str(result)
        return StructuredJSONResult(source_file=metadata.original_filename if metadata else "unknown",
                                    path=str(path), keys_involved=keys, data=data,
                                    evidence_reference=evidence_id)


_service: Optional[EvidenceService] = None


def get_evidence_service() -> EvidenceService:
    global _service
    if _service is None:
        _service = EvidenceService()
    return _service
