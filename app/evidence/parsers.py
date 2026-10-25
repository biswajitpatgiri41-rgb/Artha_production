"""File format parsing and routing for evidence files."""

from __future__ import annotations

import io
import json
import logging
import re
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Any, BinaryIO, Dict, List, Optional, Tuple

import pandas as pd
from pypdf import PdfReader

from app.core.config import get_evidence_limits
from app.evidence.schemas import (
    EvidenceDocument,
    EvidenceFormat,
    EvidenceMetadata,
    EvidenceStatus,
)

logger = logging.getLogger(__name__)
MAX_EVIDENCE_COLUMNS = 100


class ParseWarning(Exception):
    """Warning during parsing that doesn't prevent completion."""
    pass


class ParserBase(ABC):
    """Base class for file parsers."""

    @abstractmethod
    def parse(self, file_path: Path, metadata: EvidenceMetadata) -> EvidenceDocument:
        """Parse a file and return an EvidenceDocument."""
        pass

    def _add_warning(self, metadata: Optional[EvidenceMetadata], warning: str) -> None:
        """Add a warning to metadata."""
        if metadata is not None:
            metadata.parser_warnings.append(warning)
        logger.warning(warning)

    def _validate_mime_type(self, file_path: Path, expected_mime: str) -> None:
        """Compatibility no-op: upload validation inspects actual signatures and parser structure."""
        return None


class PDFParser(ParserBase):
    """Parser for PDF files."""

    def __init__(self):
        self._limits = get_evidence_limits()

    def parse(self, file_path: Path, metadata: EvidenceMetadata) -> EvidenceDocument:
        text_parts = []
        pages = []

        try:
            reader = PdfReader(str(file_path))
            page_count = len(reader.pages)
            metadata.page_count = page_count

            # Check page limit
            if page_count > self._limits["max_pdf_pages"]:
                metadata.status = EvidenceStatus.FAILED
                metadata.error_message = f"PDF has {page_count} pages, exceeds maximum of {self._limits['max_pdf_pages']}"
                raise ValueError(metadata.error_message)

            for page_num, page in enumerate(reader.pages, 1):
                try:
                    page_text = (page.extract_text() or "").replace("\x00", " ")
                    text_parts.append(page_text)
                    pages.append({"page": page_num, "text": page_text})
                except Exception as e:
                    self._add_warning(metadata, f"Failed to extract text from page {page_num}: {e}")
                    pages.append({"page": page_num, "text": ""})

        except Exception as e:
            if metadata.status != EvidenceStatus.FAILED:
                metadata.status = EvidenceStatus.FAILED
                metadata.error_message = f"PDF parsing failed: {e}"
            raise

        full_text = "\n\n".join(text_parts)

        return EvidenceDocument(
            evidence_id=metadata.evidence_id,
            run_id=metadata.run_id,
            format=EvidenceFormat.PDF,
            filename=metadata.original_filename,
            text=full_text,
            pages=pages,
        )


class TextParser(ParserBase):
    """Parser for plain text and markdown files."""

    def parse(self, file_path: Path, metadata: EvidenceMetadata) -> EvidenceDocument:
        try:
            text = file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            try:
                text = file_path.read_text(encoding="latin-1")
                self._add_warning(metadata, "File decoded as latin-1 (not UTF-8)")
            except Exception as e:
                metadata.status = EvidenceStatus.FAILED
                metadata.error_message = f"Text file decoding failed: {e}"
                raise

        sections = []
        if metadata.format == EvidenceFormat.MD:
            current_title = "Introduction"
            current_lines: list[str] = []
            for line in text.splitlines():
                heading = re.match(r"^#{1,6}\s+(.+?)\s*#*\s*$", line)
                if heading:
                    if current_lines:
                        sections.append({"page": 1, "section": current_title, "text": "\n".join(current_lines).strip()})
                    current_title = heading.group(1).strip()
                    current_lines = []
                else:
                    current_lines.append(line)
            if current_lines:
                sections.append({"page": 1, "section": current_title, "text": "\n".join(current_lines).strip()})
        else:
            sections = [{"page": 1, "section": None, "text": text}]

        return EvidenceDocument(
            evidence_id=metadata.evidence_id,
            run_id=metadata.run_id,
            format=metadata.format,
            filename=metadata.original_filename,
            text=text,
            pages=sections,
        )


class CSVParser(ParserBase):
    """Parser for CSV files."""

    def __init__(self):
        self._limits = get_evidence_limits()

    def parse(self, file_path: Path, metadata: EvidenceMetadata) -> EvidenceDocument:
        try:
            # Read only one row beyond the configured limit to bound memory.
            df = pd.read_csv(file_path, low_memory=False, nrows=self._limits["max_csv_rows"] + 1)
            if len(df) > self._limits["max_csv_rows"]:
                raise ValueError(f"CSV exceeds maximum of {self._limits['max_csv_rows']} rows")
        except Exception as e:
            metadata.status = EvidenceStatus.FAILED
            metadata.error_message = f"CSV parsing failed: {e}"
            raise

        metadata.columns = df.columns.tolist()
        if len(df.columns) > MAX_EVIDENCE_COLUMNS:
            metadata.status = EvidenceStatus.FAILED
            metadata.error_message = f"CSV exceeds the {MAX_EVIDENCE_COLUMNS} column limit"
            raise ValueError(metadata.error_message)
        metadata.row_count = len(df)
        metadata.data_types = {col: str(dtype) for col, dtype in df.dtypes.items()}

        # Convert to records for storage (limit to first 1000 rows for memory)
        records = df.head(1000).to_dict(orient="records")

        tables = [{
            "sheet": "data",
            "columns": df.columns.tolist(),
            "row_count": len(df),
            "data_types": metadata.data_types,
            "sample_data": records,
        }]

        return EvidenceDocument(
            evidence_id=metadata.evidence_id,
            run_id=metadata.run_id,
            format=EvidenceFormat.CSV,
            filename=metadata.original_filename,
            tables=tables,
        )


class XLSXParser(ParserBase):
    """Parser for Excel files (XLSX)."""

    def __init__(self):
        self._limits = get_evidence_limits()

    def parse(self, file_path: Path, metadata: EvidenceMetadata) -> EvidenceDocument:
        try:
            # Read-only, data-only workbook access never evaluates formulas or
            # executes macros. Values are consumed in bounded row batches.
            from openpyxl import load_workbook
            workbook = load_workbook(file_path, read_only=True, data_only=True, keep_links=False)
            metadata.sheet_names = list(workbook.sheetnames)
        except Exception as e:
            metadata.status = EvidenceStatus.FAILED
            metadata.error_message = f"XLSX parsing failed: {e}"
            raise

        # Check sheet limit
        if len(workbook.sheetnames) > self._limits["max_xlsx_sheets"]:
            workbook.close()
            metadata.status = EvidenceStatus.FAILED
            metadata.error_message = f"XLSX has too many sheets (max {self._limits['max_xlsx_sheets']})"
            raise ValueError(metadata.error_message)

        tables = []
        total_rows = 0
        all_columns = []

        for sheet_name in workbook.sheetnames:
            try:
                sheet = workbook[sheet_name]
                values = sheet.iter_rows(values_only=True)
                headers = next(values, ())
                columns = [str(value) if value is not None else f"column_{index + 1}" for index, value in enumerate(headers)]
                if len(columns) > MAX_EVIDENCE_COLUMNS:
                    self._add_warning(metadata, f"Sheet '{sheet_name}' has more than {MAX_EVIDENCE_COLUMNS} columns; extra columns were omitted")
                    columns = columns[:MAX_EVIDENCE_COLUMNS]
                records = []
                row_count = 0
                for raw_row in values:
                    if not any(value is not None for value in raw_row):
                        continue
                    row_count += 1
                    if row_count <= self._limits["max_csv_rows"]:
                        records.append(dict(zip(columns, raw_row[:len(columns)])))
                    if row_count == self._limits["max_csv_rows"] + 1:
                        self._add_warning(metadata, f"Sheet '{sheet_name}' exceeds the configured row limit; metadata and query sample are truncated")
                        break
                df = pd.DataFrame(records, columns=columns)
            except Exception as e:
                self._add_warning(metadata, f"Failed to parse sheet '{sheet_name}' ({type(e).__name__})")
                continue

            total_rows += row_count
            all_columns.extend(df.columns.tolist())

            sample_records = df.head(1000).to_dict(orient="records")
            tables.append({
                "sheet": sheet_name,
                "columns": df.columns.tolist(),
                "row_count": row_count,
                "data_types": {col: str(dtype) for col, dtype in df.dtypes.items()},
                "sample_data": sample_records,
            })

        workbook.close()

        metadata.row_count = total_rows
        metadata.columns = list(set(all_columns))
        metadata.data_types = {}  # Mixed across sheets

        return EvidenceDocument(
            evidence_id=metadata.evidence_id,
            run_id=metadata.run_id,
            format=EvidenceFormat.XLSX,
            filename=metadata.original_filename,
            tables=tables,
        )


class JSONParser(ParserBase):
    """Parser for JSON files."""

    def __init__(self):
        self._limits = get_evidence_limits()

    def parse(self, file_path: Path, metadata: EvidenceMetadata) -> EvidenceDocument:
        # Check file size limit for JSON
        file_size = file_path.stat().st_size
        if file_size > self._limits["max_json_bytes"]:
            metadata.status = EvidenceStatus.FAILED
            metadata.error_message = f"JSON file is {file_size} bytes, exceeds maximum of {self._limits['max_json_bytes']} bytes"
            raise ValueError(metadata.error_message)

        try:
            text = file_path.read_text(encoding="utf-8")
            data = json.loads(text)
        except json.JSONDecodeError as e:
            metadata.status = EvidenceStatus.FAILED
            metadata.error_message = f"JSON parsing failed: {e}"
            raise
        except Exception as e:
            metadata.status = EvidenceStatus.FAILED
            metadata.error_message = f"JSON file read failed: {e}"
            raise

        metadata.top_level_keys = list(data.keys()) if isinstance(data, dict) else []
        metadata.json_schema = self._infer_schema(data)

        return EvidenceDocument(
            evidence_id=metadata.evidence_id,
            run_id=metadata.run_id,
            format=EvidenceFormat.JSON,
            filename=metadata.original_filename,
            json_data=data,
            json_schema=metadata.json_schema,
        )

    def _infer_schema(self, data: Any, path: str = "") -> Dict[str, Any]:
        """Infer a simple schema from JSON data."""
        if isinstance(data, dict):
            return {k: self._infer_schema(v, f"{path}.{k}") for k, v in data.items()}
        elif isinstance(data, list):
            if data:
                return {"type": "array", "items": self._infer_schema(data[0], path)}
            return {"type": "array", "items": {}}
        elif isinstance(data, str):
            return {"type": "string"}
        elif isinstance(data, bool):
            return {"type": "boolean"}
        elif isinstance(data, (int, float)):
            return {"type": "number"}
        else:
            return {"type": "null"}


class FormatRouter:
    """Routes files to appropriate parsers based on format."""

    def __init__(self):
        self._parsers: Dict[EvidenceFormat, ParserBase] = {
            EvidenceFormat.PDF: PDFParser(),
            EvidenceFormat.CSV: CSVParser(),
            EvidenceFormat.XLSX: XLSXParser(),
            EvidenceFormat.JSON: JSONParser(),
            EvidenceFormat.TXT: TextParser(),
            EvidenceFormat.MD: TextParser(),
        }

    def parse_file(self, file_path: Path, metadata: EvidenceMetadata) -> EvidenceDocument:
        """Parse a file using the appropriate parser."""
        parser = self._parsers.get(metadata.format)
        if not parser:
            metadata.status = EvidenceStatus.FAILED
            metadata.error_message = f"No parser for format: {metadata.format}"
            raise ValueError(metadata.error_message)

        metadata.status = EvidenceStatus.PARSING
        try:
            doc = parser.parse(file_path, metadata)
            metadata.status = EvidenceStatus.READY
            return doc
        except Exception as e:
            if metadata.status != EvidenceStatus.FAILED:
                metadata.status = EvidenceStatus.FAILED
                metadata.error_message = str(e)
            raise


# Singleton router
_router: Optional[FormatRouter] = None


def get_format_router() -> FormatRouter:
    """Get the global format router instance."""
    global _router
    if _router is None:
        _router = FormatRouter()
    return _router


def parse_file(file_path: Path, metadata: EvidenceMetadata) -> EvidenceDocument:
    """Convenience function to parse a file."""
    return get_format_router().parse_file(file_path, metadata)
