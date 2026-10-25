"""Evidence retrieval for RAG."""

from __future__ import annotations

from datetime import datetime, timezone
from typing import List, Optional
import json
from app.core.config import get_evidence_limits

from app.evidence.schemas import (
    EvidenceFormat,
    RetrievedEvidence,
    RetrievedEvidenceLocation,
    SearchRequest,
    SearchResponse,
)
from app.evidence.embeddings import get_embedding_provider
from app.evidence.vector_store import get_vector_store
from app.evidence.storage import get_storage

# Cosine scores below this threshold are usually unrelated for the configured
# sentence embedding models. Returning no result is safer than passing weak
# matches to the assistant as factual support.
MIN_RELEVANCE_SCORE = 0.45

async def retrieve_evidence(request: SearchRequest) -> SearchResponse:
    """Retrieve evidence for a query within a run.

    This performs vector search for text-based documents (PDF, TXT, MD)
    and routes structured queries appropriately.
    """
    storage = get_storage()
    vector_store = get_vector_store()

    metadata_by_id = {item.evidence_id: item for item in storage.list_metadata(request.run_id)}
    if request.evidence_ids is not None:
        allowed_ids = set(request.evidence_ids)
        metadata_by_id = {key: value for key, value in metadata_by_id.items() if key in allowed_ids}

    # For now, perform vector search on the run's collection
    text_formats = {EvidenceFormat.PDF, EvidenceFormat.TXT, EvidenceFormat.MD}
    has_ready_text = any(item.status.value == "ready" and item.format in text_formats for item in metadata_by_id.values())
    matches = []
    if has_ready_text:
        embedding_provider = get_embedding_provider()
        embedding_model = embedding_provider.model_name
        query_embedding = embedding_provider.embed_single(request.query)
        # Oversample before relevance filtering so a few weak/duplicate chunks do
        # not hide a useful match just below the initial top-k boundary.
        candidate_count = min(max(request.top_k * 4, request.top_k), 50)
        matches = vector_store.search(request.run_id, query_embedding, candidate_count,
                                      request.evidence_ids)

    results = []
    seen_chunks = set()
    for match in matches:
        threshold = 0.20 if has_ready_text and embedding_model == "local-hash-v1" else MIN_RELEVANCE_SCORE
        if match["score"] < threshold:
            continue
        meta = match["metadata"]
        # The collection is isolated by run, and metadata is checked again so
        # stale or malformed vector records cannot cross the run boundary.
        if meta.get("run_id") != request.run_id:
            continue
        file_metadata = metadata_by_id.get(meta.get("evidence_id", ""))
        if file_metadata is None or file_metadata.status.value != "ready":
            continue
        if any("instruction-like text detected" in warning.casefold() for warning in file_metadata.parser_warnings):
            # Do not forward passages from a document flagged during full-text
            # parsing; this also catches injection strings split across chunks.
            continue
        if has_ready_text and embedding_model == "local-hash-v1":
            query_words = {word.casefold() for word in request.query.split() if len(word) > 2}
            chunk_words = {word.casefold() for word in match["text"].split()}
            if not query_words.intersection(chunk_words):
                continue
        # Overlapping PDF chunks can repeat the same passage. Keep one copy.
        fingerprint = (meta.get("evidence_id"), " ".join(match["text"].split()).casefold())
        if fingerprint in seen_chunks:
            continue
        seen_chunks.add(fingerprint)
        location = RetrievedEvidenceLocation(
            page=meta.get("page"),
            section=meta.get("section"),
            chunk_id=match["chunk_id"],
        )
        evidence = RetrievedEvidence(
            evidence_id=meta["evidence_id"],
            filename=meta["filename"],
            format=file_metadata.format,
            location=location,
            excerpt=match["text"][:500],
            retrieval_score=match["score"],
            run_id=request.run_id,
            retrieved_at=datetime.now(timezone.utc),
        )
        results.append(evidence)
        if len(results) >= request.top_k:
            break

    warnings = []
    # Structured files use deterministic, bounded text matching. They are
    # never sent through vector similarity or expanded into embedding text.
    query_terms = {term.casefold() for term in request.query.split() if len(term) > 2}
    structured_count = 0
    for file_metadata in metadata_by_id.values():
        if len(results) >= request.top_k or structured_count >= request.top_k or file_metadata.status.value != "ready":
            continue
        if file_metadata.format not in (EvidenceFormat.CSV, EvidenceFormat.XLSX, EvidenceFormat.JSON):
            continue
        if file_metadata.format == EvidenceFormat.JSON:
            keys = [str(key)[:80] for key in (file_metadata.top_level_keys or [])[:30]]
            matching = [key for key in keys if any(term in key.casefold() for term in query_terms)]
            if matching:
                results.append(RetrievedEvidence(
                    evidence_id=file_metadata.evidence_id, filename=file_metadata.original_filename,
                    format=file_metadata.format,
                    location=RetrievedEvidenceLocation(keys=matching, path="$"),
                    structured_result={"top_level_keys": keys, "matching_keys": matching},
                    run_id=request.run_id, retrieved_at=datetime.now(timezone.utc),
                ))
                structured_count += 1
            continue
        file_path = storage.get_file_path(request.run_id, file_metadata.stored_filename)
        import pandas as pd
        if file_metadata.format == EvidenceFormat.CSV:
            frame = pd.read_csv(file_path, nrows=get_evidence_limits()["max_csv_rows"] + 1, low_memory=False)
            sheet_name = None
        else:
            from openpyxl import load_workbook
            sheet_name = (file_metadata.sheet_names or ["Sheet1"])[0]
            workbook = load_workbook(file_path, read_only=True, data_only=True, keep_links=False)
            try:
                worksheet = workbook[sheet_name]
                iterator = worksheet.iter_rows(values_only=True)
                header = next(iterator, ())
                columns = [str(value) if value is not None else f"column_{index + 1}" for index, value in enumerate(header)]
                rows = []
                for row in iterator:
                    if len(rows) >= get_evidence_limits()["max_csv_rows"]:
                        break
                    rows.append(row)
                frame = pd.DataFrame(rows, columns=columns)
            finally:
                workbook.close()
        if not query_terms:
            continue
        matching_columns = [str(col) for col in frame.columns if any(term in str(col).casefold() for term in query_terms)]
        if query_terms:
            as_text = frame.astype(str).apply(lambda column: column.str.casefold())
            row_mask = pd.Series(False, index=frame.index)
            for term in query_terms:
                row_mask = row_mask | as_text.apply(lambda column: column.str.contains(term, regex=False)).any(axis=1)
            matching_rows = frame.loc[row_mask].head(10)
        else:
            matching_rows = frame.iloc[0:0]
        if matching_columns or not matching_rows.empty:
            selected_columns = (matching_columns or [str(column) for column in matching_rows.columns])[:8]
            sample = matching_rows[selected_columns].head(3)
            sample_records = json.loads(sample.to_json(orient="records", date_format="iso"))
            for record in sample_records:
                for key, value in record.items():
                    if isinstance(value, str) and len(value) > 120:
                        record[key] = value[:120] + "…"
            results.append(RetrievedEvidence(
                evidence_id=file_metadata.evidence_id, filename=file_metadata.original_filename,
                format=file_metadata.format,
                location=RetrievedEvidenceLocation(rows=[int(index) + 2 for index in sample.index], columns=selected_columns, sheet=sheet_name),
                structured_result={"matching_columns": matching_columns, "schema": {column: str(frame[column].dtype) for column in selected_columns if column in frame.columns},
                                   "rows": sample_records,
                                   "row_count": file_metadata.row_count},
                run_id=request.run_id, retrieved_at=datetime.now(timezone.utc),
            ))
            structured_count += 1

    if not results:
        warnings.append(
            "No sufficiently relevant evidence found. Retrieved passages were below the relevance threshold."
        )

    return SearchResponse(results=results, warnings=warnings)


def retrieve_structured_evidence(
    run_id: str,
    evidence_id: str,
    query: str,
) -> Optional[List[RetrievedEvidence]]:
    """Retrieve from structured evidence (CSV, XLSX, JSON).

    This is a placeholder for the structured evidence tools.
    """
    # TODO: Implement structured evidence retrieval
    # This would use the structured evidence tools from service.py
    return None
